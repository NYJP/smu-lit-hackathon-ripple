"""Real upload coverage for the stage-three ingestion boundary."""

from __future__ import annotations

import fitz
from docx import Document as DocxDocument

from api.services import parsing, storage
from tests.conftest import login_as


def _pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    result = document.tobytes()
    document.close()
    return result


def test_regulation_pdf_is_ingested_ready_with_no_requirements(client, conn, users):
    login_as(client, users["Priya Menon"]["id"])
    response = client.post(
        "/api/v1/regulations",
        data={"title": "Retention Rules", "document_kind": "primary"},
        files={"file": ("retention.pdf", _pdf_bytes("Section 1\nRecords must be retained for seven years."), "application/pdf")},
    )
    assert response.status_code == 202, response.text
    created = response.json()
    job = client.get(f"/api/v1/jobs/{created['job_id']}")
    assert job.status_code == 200
    body = job.json()
    assert {key: body[key] for key in ("id", "job_type", "status", "progress", "step", "error_message")} == {
        "id": created["job_id"], "job_type": "regulation_ingest", "status": "succeeded",
        "progress": 1.0, "step": "Ready", "error_message": None,
    }
    assert body["result"]["page_count"] == 1
    assert body["result"]["requirement_count"] == 0
    assert body["result"]["token_usage"]["total_tokens"] == 2
    assert body["result"]["estimated_cost_usd"] > 0
    regulation = conn.execute("SELECT status, page_count FROM regulations WHERE id = ?", (created["regulation_id"],)).fetchone()
    assert dict(regulation) == {"status": "ready", "page_count": 1}
    assert conn.execute("SELECT COUNT(*) AS c FROM regulatory_requirements").fetchone()["c"] == 0
    file_response = client.get(f"/api/v1/files/regulations/{created['regulation_id']}")
    assert file_response.status_code == 200
    assert file_response.headers["content-type"].startswith("application/pdf")


def test_document_docx_ingestion_writes_chunks_fts_and_scopes_collaborator(client, conn, users, tmp_path):
    docx = DocxDocument()
    docx.add_heading("Part I Retention", level=1)
    docx.add_paragraph("Customer records are retained for seven years in the approved archive.")
    docx_path = tmp_path / "policy.docx"
    docx.save(docx_path)
    login_as(client, users["Alex Tan"]["id"])
    response = client.post(
        "/api/v1/documents",
        data={"doc_type": "policy", "collaborator_ids[]": users["Sam Rahim"]["id"], "collaborator_access": "reviewer"},
        files=[("files[]", ("retention-policy.docx", docx_path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert response.status_code == 202, response.text
    item = response.json()["documents"][0]
    assert item["scan_id"] is None
    job = client.get(f"/api/v1/jobs/{item['job_id']}").json()
    assert job["status"] == "succeeded"
    assert job["result"]["chunk_count"] == 2
    chunks = conn.execute("SELECT content, section_path, chunk_type FROM document_chunks WHERE document_id = ? ORDER BY ordinal", (item["document_id"],)).fetchall()
    assert [row["chunk_type"] for row in chunks] == ["heading", "paragraph"]
    assert chunks[1]["section_path"] == "Part I Retention"
    fts = conn.execute("SELECT content FROM fts_chunks WHERE chunk_id IN (SELECT id FROM document_chunks WHERE document_id = ?)", (item["document_id"],)).fetchall()
    assert any("seven years" in row["content"] for row in fts)

    login_as(client, users["Sam Rahim"]["id"])
    shared = client.get("/api/v1/documents?scope=shared_with_me")
    assert shared.status_code == 200
    assert [row["id"] for row in shared.json()["items"]] == [item["document_id"]]
    detail = client.get(f"/api/v1/documents/{item['document_id']}")
    assert detail.status_code == 200
    assert len(detail.json()["chunks"]) == 2
    streamed = client.get(f"/api/v1/files/documents/{item['document_id']}")
    assert streamed.status_code == 200
    assert streamed.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def test_document_pdf_multpage_sections_keep_page_and_source_provenance(client, conn, users):
    """Generated pages retain their own heading, path, and original offset range."""
    pdf = fitz.open()
    retention_page = pdf.new_page()
    retention_page.insert_text((72, 72), "Part I Retention")
    retention_page.insert_text((72, 88), "Customer records are retained for seven years.")
    access_page = pdf.new_page()
    access_page.insert_text((72, 72), "Part II Access")
    access_page.insert_text((72, 88), "Access logs are retained for seven years.")
    pdf_payload = pdf.tobytes()
    pdf.close()
    login_as(client, users["Alex Tan"]["id"])
    response = client.post(
        "/api/v1/documents",
        data={"doc_type": "policy"},
        files=[("files[]", ("retention.pdf", pdf_payload, "application/pdf"))],
    )
    assert response.status_code == 202, response.text
    document_id = response.json()["documents"][0]["document_id"]
    job = client.get(f"/api/v1/jobs/{response.json()['documents'][0]['job_id']}").json()
    assert job["status"] == "succeeded", job
    chunks = conn.execute(
        "SELECT content, section_path, page_number, char_start, char_end, chunk_type "
        "FROM document_chunks WHERE document_id = ? ORDER BY ordinal",
        (document_id,),
    ).fetchall()
    assert [chunk["chunk_type"] for chunk in chunks] == ["heading", "paragraph", "heading", "paragraph"]
    assert [chunk["page_number"] for chunk in chunks] == [1, 1, 2, 2]
    assert [chunk["section_path"] for chunk in chunks] == [None, "Part I Retention", "Part I Retention", "Part II Access"]

    stored = conn.execute("SELECT file_path FROM documents WHERE id = ?", (document_id,)).fetchone()
    original_extracted_text = "\n".join(parsing._extract_raw_pages(storage.absolute_path(stored["file_path"])))
    for chunk in chunks:
        assert original_extracted_text[chunk["char_start"]:chunk["char_end"]] == chunk["content"]


def test_document_txt_offsets_preserve_original_formatting(client, conn, users):
    # The body deliberately has a newline and repeated spaces. The former
    # normalized/reconstructed offset scheme could not make this source slice
    # equal the stored content.
    source_text = "Part II Access\n\nAccess logs are retained   for seven years,\nwith quarterly verification."
    login_as(client, users["Alex Tan"]["id"])
    response = client.post(
        "/api/v1/documents",
        data={"doc_type": "policy"},
        files=[("files[]", ("access.txt", source_text.encode(), "text/plain"))],
    )
    assert response.status_code == 202, response.text
    document_id = response.json()["documents"][0]["document_id"]
    chunks = conn.execute(
        "SELECT content, section_path, char_start, char_end, chunk_type "
        "FROM document_chunks WHERE document_id = ? ORDER BY ordinal",
        (document_id,),
    ).fetchall()
    assert [chunk["chunk_type"] for chunk in chunks] == ["heading", "paragraph"]
    assert chunks[1]["section_path"] == "Part II Access"
    assert "   " in chunks[1]["content"]
    assert "\n" in chunks[1]["content"]
    for chunk in chunks:
        assert source_text[chunk["char_start"]:chunk["char_end"]] == chunk["content"]


def test_scanned_pdf_fails_with_the_required_ocr_message(client, conn, users):
    login_as(client, users["Alex Tan"]["id"])
    # A blank PDF yields no extractable text and therefore crosses the 30% threshold.
    blank = fitz.open()
    blank.new_page()
    payload = blank.tobytes()
    blank.close()
    response = client.post(
        "/api/v1/documents", data={"doc_type": "policy"},
        files=[("files[]", ("scan.pdf", payload, "application/pdf"))],
    )
    assert response.status_code == 202
    item = response.json()["documents"][0]
    job = client.get(f"/api/v1/jobs/{item['job_id']}").json()
    assert job["status"] == "failed"
    assert job["error_message"] == "Scanned PDF — OCR is not supported in the MVP"
    document = conn.execute("SELECT status, error_message FROM documents WHERE id = ?", (item["document_id"],)).fetchone()
    assert dict(document) == {"status": "failed", "error_message": "Scanned PDF — OCR is not supported in the MVP"}
